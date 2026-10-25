import io

import docx
import httpx
import pytest

from app.errors import PermanentError
from app.models import Item, SourceType
from app.sources.files import FileSource, paragraphs_from_plain_text
from app.sources.note import NoteSource
from app.sources.youtube import YoutubeSource, timestamp, video_id


def _file_item(tmp_path, name: str, data: bytes) -> Item:
    path = tmp_path / name
    path.write_bytes(data)
    return Item(title=name, source_type=SourceType.file, source=str(path))


def test_plain_text_paragraphs_join_wrapped_lines_but_keep_lists():
    text = "First line\nwrapped here.\n\n- one\n- two\n\n1. a\n2. b"
    assert paragraphs_from_plain_text(text) == [
        "First line wrapped here.",
        "- one",
        "- two",
        "1. a",
        "2. b",
    ]


def test_note_source_reads_the_text_on_the_item():
    doc = NoteSource().extract(Item(source_type=SourceType.note, source="Para one.\n\nPara two."))
    assert doc.pages[0].text == "Para one.\n\nPara two."


def test_markdown_loses_its_syntax_and_gives_a_title(tmp_path):
    md = (
        b"# Hash tables\n\nUse **open addressing** or [chaining](https://x.com).\n\n"
        b"## Notes\n\n`O(1)` average."
    )
    doc = FileSource().extract(_file_item(tmp_path, "notes.md", md))
    assert doc.title == "Hash tables"
    text = doc.pages[0].text
    assert "Use open addressing or chaining." in text and "O(1) average." in text
    assert "#" not in text and "**" not in text and "https" not in text


def test_word_documents_include_tables(tmp_path):
    document = docx.Document()
    document.core_properties.title = "Lecture 4"
    document.add_paragraph("B-trees keep keys sorted.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Index"
    table.rows[0].cells[1].text = "B-tree"
    buffer = io.BytesIO()
    document.save(buffer)
    doc = FileSource().extract(_file_item(tmp_path, "l4.docx", buffer.getvalue()))
    assert doc.title == "Lecture 4"
    assert doc.pages[0].text == "B-trees keep keys sorted.\n\nIndex | B-tree"


def test_damaged_word_file_is_a_permanent_error(tmp_path):
    with pytest.raises(PermanentError, match="Word document"):
        FileSource().extract(_file_item(tmp_path, "bad.docx", b"not a zip"))


@pytest.mark.parametrize(
    ("url", "expected"),
    [
        ("https://www.youtube.com/watch?v=kCc8FmEb1nY&t=10s", "kCc8FmEb1nY"),
        ("https://youtu.be/kCc8FmEb1nY", "kCc8FmEb1nY"),
        ("https://youtube.com/shorts/kCc8FmEb1nY", "kCc8FmEb1nY"),
        ("https://m.youtube.com/watch?v=kCc8FmEb1nY", "kCc8FmEb1nY"),
        ("https://youtube.com/about", None),
        ("https://example.com/watch?v=kCc8FmEb1nY", None),
    ],
)
def test_video_ids(url, expected):
    assert video_id(url) == expected


def test_timestamps():
    assert timestamp(5) == "0:05" and timestamp(754) == "12:34" and timestamp(3725) == "1:02:05"


def test_youtube_transcript_becomes_timed_paragraphs():
    snippets = [(0.0, "hello everyone"), (30.0, "[Music]"), (45.0, "today we build"), (70.0, "GPT")]
    oembed = httpx.Client(
        transport=httpx.MockTransport(lambda req: httpx.Response(200, json={"title": "Build GPT"}))
    )
    source = YoutubeSource(fetch=lambda vid: snippets, client=oembed)
    doc = source.extract(
        Item(source_type=SourceType.youtube, source="https://youtu.be/kCc8FmEb1nY")
    )
    assert doc.title == "Build GPT"
    assert doc.pages[0].text == "[0:00] hello everyone today we build GPT"


def test_youtube_without_captions_is_a_permanent_error():
    def no_captions(vid):
        raise PermanentError("This video has no captions, so there's nothing to read.")

    source = YoutubeSource(
        fetch=no_captions,
        client=httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(404))),
    )
    with pytest.raises(PermanentError, match="no captions"):
        source.extract(Item(source_type=SourceType.youtube, source="https://youtu.be/kCc8FmEb1nY"))
